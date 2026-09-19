"""Rich DOCX/PDF writers used by the existing export services."""

from io import BytesIO
from xml.sax.saxutils import escape

from docx.shared import Inches
from PIL import Image as PILImage
from reportlab.platypus import Flowable, Image, Table, TableStyle, Paragraph
from reportlab.lib import colors

from app.services.math_rendering import OfficeMathConverter, formula_png, math_segments, math_error
from app.services.rich_content import diagram_png, image_bytes, normalize_blocks


def _image_size(data, max_width=450, max_height=500):
    with PILImage.open(BytesIO(data)) as image:
        width, height = image.size
    scale = min(max_width / width, max_height / height, 1)
    return width * scale, height * scale


def docx_text(parent, text: str):
    paragraph = parent.add_paragraph()
    try:
        for kind, value, display in math_segments(str(text)):
            if kind == "text":
                paragraph.add_run(value)
            else:
                try:
                    equation = OfficeMathConverter().omml(value)
                    if display and paragraph.text:
                        paragraph = parent.add_paragraph()
                    paragraph._p.append(equation)
                    if display:
                        paragraph = parent.add_paragraph()
                except Exception:
                    data = formula_png(value)
                    width, height = _image_size(data, max_width=430, max_height=100)
                    paragraph.add_run().add_picture(BytesIO(data), width=Inches(width / 120), height=Inches(height / 120))
    except Exception as error:
        raise math_error(error) from error
    return paragraph


def docx_blocks(parent, value):
    for block in normalize_blocks(value):
        kind = block["type"]
        if kind == "text":
            docx_text(parent, block["content"])
        elif kind == "latex":
            delimiter = "$$" if block["display"] == "block" else "$"
            docx_text(parent, delimiter + block["content"] + delimiter)
        elif kind == "table":
            table = parent.add_table(rows=1, cols=len(block["headers"]))
            table.style = "Table Grid"
            for index, cell in enumerate(block["headers"]):
                docx_text(table.rows[0].cells[index], cell)
            for row in block["rows"]:
                cells = table.add_row().cells
                for index, cell in enumerate(row):
                    docx_text(cells[index], cell)
        else:
            data = diagram_png(block["spec"]) if kind == "diagram" else image_bytes(block["src"])
            width, height = _image_size(data)
            run = parent.add_paragraph().add_run()
            inline = run.add_picture(BytesIO(data), width=Inches(width / 80), height=Inches(height / 80))
            inline._inline.docPr.set("descr", block["alt"])
        if block.get("caption"):
            docx_text(parent, block["caption"])


class FormulaFlowable(Flowable):
    def __init__(self, expression):
        super().__init__()
        try:
            self.data = formula_png(expression)
        except Exception as error:
            raise math_error(error) from error
        self.natural_width, self.natural_height = _image_size(self.data, 430, 120)
        # 180 dpi raster -> points, retain readability at normal print scale.
        self.natural_width *= 72 / 180
        self.natural_height *= 72 / 180

    def wrap(self, available_width, available_height):
        scale = min(1, available_width / max(1, self.natural_width))
        self.width, self.height = self.natural_width * scale, self.natural_height * scale
        return self.width, self.height

    def draw(self):
        from reportlab.lib.utils import ImageReader
        self.canv.drawImage(ImageReader(BytesIO(self.data)), 0, 0, self.width, self.height, mask="auto")


def pdf_text(text, style):
    """Flowables preserve text wrapping; formula segments are separate safe runs."""
    from xml.sax.saxutils import escape
    result = []
    for kind, value, _display in math_segments(str(text)):
        if kind == "latex":
            result.append(FormulaFlowable(value))
        elif value.strip():
            result.append(Paragraph(escape(value).replace("\n", "<br/>"), style))
    return result or [Paragraph("", style)]


def pdf_paragraph(text, style):
    # Build image fragments from in-memory bytes. ReportLab URL loading remains
    # disabled, including data URLs; no global trust settings are changed.
    from copy import copy
    from reportlab.lib.abag import ABag
    from reportlab.lib.utils import ImageReader
    fragments = []
    try:
        for kind, value, display in math_segments(str(text)):
            if kind == "text":
                escaped = escape(value).replace("\n", "<br/>")
                if value.startswith(" "):
                    escaped = "&nbsp;" + escaped[1:]
                if value.endswith(" ") and len(value) > 1:
                    escaped = escaped[:-1] + "&nbsp;"
                fragments.extend(Paragraph(escaped, style).frags)
            else:
                data = formula_png(value)
                width, height = _image_size(data, max_width=1000, max_height=300)
                width, height = width * 72 / 180, height * 72 / 180
                if display:
                    fragments.extend(Paragraph("<br/>", style).frags)
                fragment = copy(Paragraph("x", style).frags[0])
                fragment.text = ""
                fragment.cbDefn = ABag(kind="img", image=ImageReader(BytesIO(data)), width=width, height=height, valign="middle", src="embedded-formula")
                fragments.append(fragment)
                if display:
                    fragments.extend(Paragraph("<br/>", style).frags)
    except Exception as error:
        raise math_error(error) from error
    return Paragraph("", style, frags=fragments)


def pdf_blocks(value, style):
    result = []
    for block in normalize_blocks(value):
        kind = block["type"]
        if kind == "text":
            result.extend(pdf_text(block["content"], style))
        elif kind == "latex":
            result.append(FormulaFlowable(block["content"]))
        elif kind == "table":
            table = Table([[pdf_text(cell, style) for cell in row] for row in [block["headers"], *block["rows"]]], colWidths=[450 / len(block["headers"])] * len(block["headers"]), repeatRows=1, splitInRow=1)
            table.setStyle(TableStyle([("GRID", (0, 0), (-1, -1), 0.5, colors.grey), ("VALIGN", (0, 0), (-1, -1), "TOP")]))
            result.append(table)
        else:
            data = diagram_png(block["spec"]) if kind == "diagram" else image_bytes(block["src"])
            width, height = _image_size(data)
            result.append(Image(BytesIO(data), width=width, height=height))
        if block.get("caption"):
            result.extend(pdf_text(block["caption"], style))
    return result
