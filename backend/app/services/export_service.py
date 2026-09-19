from docx import Document
from app.services.rich_export import docx_text, docx_blocks

from app.services.exam.export_document import DOCUMENT_TITLES, require_export_document, select_export_data


class ExportService:
    """Create one ready-to-print Word document for the selected paper."""

    DIFFICULTY_LABELS = {
        "nhan_biet": "Nhận biết",
        "thong_hieu": "Thông hiểu",
        "van_dung": "Vận dụng",
    }
    QUESTION_TYPE_LABELS = {
        "multiple_choice": "Trắc nghiệm",
        "true_false": "Đúng/Sai",
        "short_answer": "Trả lời ngắn",
        "essay": "Tự luận",
    }

    def create_full_docx(self, exam_data: dict, *, include_answers: bool = True,
                         document: str = "exam", variant_code: str | None = None) -> Document:
        require_export_document(document, include_answers)
        data = select_export_data(exam_data, variant_code)
        doc = Document()
        info = data["exam_info"]
        doc.add_heading(DOCUMENT_TITLES[document], level=0)
        docx_text(doc, f"Trường: {info['school']}")
        docx_text(doc, f"Môn: {info['subject']} - Lớp {info['grade']}")
        docx_text(doc, f"Loại kiểm tra: {info['exam_type']}")
        docx_text(doc, f"Năm học: {info['school_year']} - Thời gian: {info['duration_minutes']} phút")
        docx_text(doc, f"Mã đề {variant_code}" if variant_code else "Bản gốc")
        if document == "exam":
            self._questions(doc, data["questions"])
        elif document == "matrix":
            self._matrix(doc, data["matrix"], data["summary"])
        elif document == "specification":
            self._specification(doc, data["specification"])
        else:
            self._answers(doc, data["answer_key"])
            self._rubric(doc, data["rubric"])
        return doc

    def _matrix(self, doc: Document, matrix, summary):
        table = doc.add_table(rows=1, cols=6)
        table.style = "Table Grid"
        headers = ["Nội dung", "Nhận biết", "Thông hiểu", "Vận dụng", "Tổng câu", "Tổng điểm"]
        for index, header in enumerate(headers):
            table.rows[0].cells[index].text = header
        for row in matrix:
            cells = table.add_row().cells
            cells[0].text = row["lesson_name"]
            cells[1].text = self._cell_text(row["nhan_biet"])
            cells[2].text = self._cell_text(row["thong_hieu"])
            cells[3].text = self._cell_text(row["van_dung"])
            cells[4].text = str(row["nhan_biet"]["count"] + row["thong_hieu"]["count"] + row["van_dung"]["count"])
            cells[5].text = str(row["total_score"])
        docx_text(doc, f"Tổng điểm: {summary['total_score']} - Tổng số câu: {summary['total_questions']}")

    def _specification(self, doc: Document, specification):
        table = doc.add_table(rows=1, cols=6)
        table.style = "Table Grid"
        headers = ["Câu", "Nội dung", "Yêu cầu cần đạt", "Mức độ", "Dạng câu", "Điểm"]
        for index, header in enumerate(headers):
            table.rows[0].cells[index].text = header
        for spec in specification:
            cells = table.add_row().cells
            cells[0].text = str(spec["question_number"])
            cells[1].text = spec["knowledge_unit"]
            cells[2].text = spec["achievement"]
            cells[3].text = self._difficulty_text(spec["difficulty"])
            cells[4].text = self._question_type_text(spec["question_type"])
            cells[5].text = str(spec["score"])

    def _questions(self, doc: Document, questions, *, heading: str | None = None):
        if heading:
            doc.add_page_break()
            doc.add_heading(heading, level=1)
        for question in questions:
            docx_text(doc, f"Câu {question['number']} ({question['score']} điểm). {question['content']}")
            docx_blocks(doc, question.get("rich_content"))
            if question["type"] == "multiple_choice":
                for key, value in question.get("options", {}).items():
                    docx_text(doc, f"{key}. {value}")
            if question["type"] == "true_false":
                for statement in question.get("statements", []):
                    docx_text(doc, f"- {statement.get('content', '')}")
            if question["type"] == "essay":
                for sub_question in question.get("sub_questions") or []:
                    docx_text(doc, f"- {sub_question['content']} ({sub_question['score']} điểm)")

    def _answers(self, doc: Document, answer_key, *, heading: str | None = None):
        if heading:
            doc.add_page_break()
            doc.add_heading(heading, level=1)
        for answer in answer_key:
            prefix = f"Câu {answer.get('question_number', '')}: "
            answer_type = answer.get("type")
            if answer_type == "multiple_choice":
                docx_text(doc, prefix + (answer.get("correct_answer") or ""))
            elif answer_type == "true_false":
                values = ["Đ" if item.get("is_true") else "S" for item in answer.get("answers", [])]
                docx_text(doc, prefix + ", ".join(values))
            elif answer_type == "short_answer":
                docx_text(doc, prefix + (answer.get("correct_answer") or ""))
            else:
                docx_text(doc, prefix + (answer.get("model_answer") or "Xem hướng dẫn chấm."))

    def _rubric(self, doc: Document, rubric):
        doc.add_page_break()
        doc.add_heading("Hướng dẫn chấm", level=1)
        for item in rubric:
            docx_text(doc, f"Câu {item['question_number']} ({item['total_score']} điểm)")
            for criterion in item["criteria"]:
                docx_text(doc, f"{criterion['name']}: tối đa {criterion['max_score']} điểm")
                for level in criterion["levels"]:
                    docx_text(doc, f"- {level['score']} điểm: {level['description']}")

    def _cell_text(self, cell):
        if cell["count"] == 0:
            return "-"
        return f"{cell['count']} câu / {cell['score']} điểm"

    def _difficulty_text(self, value: str) -> str:
        return self.DIFFICULTY_LABELS.get(value, value)

    def _question_type_text(self, value: str) -> str:
        return self.QUESTION_TYPE_LABELS.get(value, value)
