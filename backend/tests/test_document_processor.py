from io import BytesIO
from types import SimpleNamespace
import os
import shutil

import pytest
from PIL import Image, ImageDraw, ImageFont
from reportlab.pdfgen import canvas
from reportlab.lib.utils import ImageReader

from app.core.config import settings
from app.extractors.document_processor import DocumentProcessor, extraction_quality, structured_blocks, normalize_content
from app.extractors.document_extractors import extract_pdf, InvalidDocumentError
from app.extractors.ocr import OCRProvider, OCRResult, TesseractOCRProvider, create_ocr_provider
from app.rag.chunking import chunk_document
from app.rag.retriever import BM25Retriever


class StubOCR(OCRProvider):
    def __init__(self, text=r'Vận tốc được tính bằng $v=\frac{s}{t}$.', fail=False):
        self.calls = []
        self.text, self.fail = text, fail
    def recognize(self, image, *, mime_type='image/png'):
        self.calls.append(image)
        assert image.startswith(b'\x89PNG')
        if self.fail:
            raise TimeoutError()
        return OCRResult(self.text, 'test', True)


def pdf(*page_texts):
    buffer = BytesIO()
    doc = canvas.Canvas(buffer)
    for text in page_texts:
        doc.drawString(45, 700, text)
        doc.showPage()
    doc.save()
    return buffer.getvalue()


def test_native_pdf_good_text_never_calls_ocr():
    data = pdf('The leaf contains chlorophyll and uses light energy in photosynthesis to produce glucose.')
    assert extract_pdf(data)['parsed_data']['pages'][0]['page'] == 1
    ocr = StubOCR()
    result = DocumentProcessor(ocr=ocr).process('text.pdf', data)
    assert ocr.calls == []
    assert result['parsed_data']['blocks'][0]['page_number'] == 1
    assert result['parsed_data']['page_quality'][0]['method'] == 'native'


def test_selective_ocr_and_index_preserve_source_page():
    data = pdf('Photosynthesis uses carbon dioxide and water and light to produce glucose in the leaves.', '')
    ocr = StubOCR('## Chuyển động\n\n$$\nv=\\frac{s}{t}\n$$\n\n| t | s |\n| 20 | 100 |')
    result = DocumentProcessor(ocr=ocr).process('mixed.pdf', data)
    assert len(ocr.calls) == 1
    chunks = chunk_document({'id': 17, 'filename': 'mixed.pdf', 'file_type': 'pdf', 'extracted_text': result['text'], 'parsed_data': result['parsed_data'], 'grade': 7}, chunk_size=15)
    formula = next(c for c in chunks if c['content_type'] == 'formula')
    assert formula['page_number'] == formula['metadata']['source_page'] == 2
    assert formula['document_id'] == 17
    assert r'\frac{s}{t}' in formula['text']
    assert len([c for c in chunks if c['content_type'] == 'table']) == 1
    hits = BM25Retriever(chunks).retrieve('Chuyển động', k=2)
    assert hits[0]['chunk']['metadata']['doc_id'] == 17


def test_ocr_normalization_preserves_formulas_and_structure():
    normalized = normalize_content('\ufeff# Va\u0323\u0302t ly\u0301\r\n\r\n$$v=\\frac{s}{t}$$\x00')
    assert '\r' not in normalized and '\x00' not in normalized
    blocks = structured_blocks(normalized, page=4)
    assert [b['content_type'] for b in blocks] == ['heading', 'formula']
    assert blocks[1]['normalized_text'] == r'$$v=\frac{s}{t}$$'


@pytest.mark.parametrize('text,needed', [('a', True), ('(cid:15) (cid:12)', True), ('+++++++ =========== ------ ########## ======', True), ('Photosynthesis uses light energy to make glucose from carbon dioxide and water.', False)])
def test_ocr_quality_decision(text, needed):
    assert extraction_quality(text)['needs_ocr'] is needed


def test_unconfigured_or_failed_ocr_retains_native_but_rejects_empty_scan():
    for ocr in [None, StubOCR(fail=True)]:
        result = DocumentProcessor(ocr=ocr).process('poor.pdf', pdf('Short caption'))
        assert result['text'] == 'Short caption'
        assert result['parsed_data']['extraction_warnings']
        with pytest.raises(InvalidDocumentError, match='OCR'):
            DocumentProcessor(ocr=ocr).process('scan.pdf', pdf(''))


def test_page_budget_and_explicit_native_mode(monkeypatch):
    monkeypatch.setattr(settings, 'OCR_MAX_PAGES', 1)
    ocr = StubOCR()
    result = DocumentProcessor(ocr=ocr).process('many.pdf', pdf('', '', ''))
    assert len(ocr.calls) == 1
    assert len(result['parsed_data']['extraction_warnings']) == 2
    DocumentProcessor(ocr=ocr).process('native.pdf', pdf('short'), ocr_mode='native')
    assert len(ocr.calls) == 1


def test_image_upload_uses_ocr_and_validates_actual_image():
    buffer = BytesIO()
    Image.new('RGB', (80, 50), 'white').save(buffer, format='PNG')
    assert DocumentProcessor(ocr=StubOCR()).process('scan.png', buffer.getvalue())['parsed_data']['ocr_pages'] == 1
    with pytest.raises(InvalidDocumentError):
        DocumentProcessor(ocr=StubOCR()).process('fake.png', b'not an image')


def test_provider_factory_gracefully_disables_missing_provider(monkeypatch):
    monkeypatch.setattr(settings, 'OCR_PROVIDER', 'gemini')
    assert create_ocr_provider() is None
    monkeypatch.setattr(settings, 'OCR_PROVIDER', 'tesseract')
    monkeypatch.setattr(settings, 'OCR_TESSERACT_COMMAND', '/does/not/exist')
    assert create_ocr_provider() is None


def test_real_tesseract_scanned_pdf(monkeypatch):
    executable = os.environ.get('TEST_TESSERACT_COMMAND') or shutil.which('tesseract')
    if not executable:
        pytest.skip('Local Tesseract executable is optional outside the backend Docker image')
    monkeypatch.setattr(settings, 'OCR_TESSERACT_COMMAND', executable)
    monkeypatch.setattr(settings, 'OCR_LANGUAGES', 'vie+eng')
    image = Image.new('RGB', (1200, 350), 'white')
    font = ImageFont.truetype(str(__import__('pathlib').Path(__file__).resolve().parents[1] / 'app/assets/fonts/DejaVuSans.ttf'), 45)
    ImageDraw.Draw(image).text((35, 45), 'Vận tốc = quãng đường / thời gian\n100 m trong 20 s. Tính vận tốc.', fill='black', font=font)
    output = BytesIO()
    doc = canvas.Canvas(output, pagesize=(600, 200))
    doc.drawImage(ImageReader(image), 0, 10, 600, 175)
    doc.save()
    result = DocumentProcessor(ocr=TesseractOCRProvider()).process('real-scan.pdf', output.getvalue())
    assert '100' in result['text'] and '20' in result['text']
    assert 'vận tốc' in result['text'].lower()
    assert result['parsed_data']['page_quality'][0]['method'] == 'ocr'
    assert result['parsed_data']['blocks'][0]['page_number'] == 1


def test_gemini_ocr_sends_image_bytes_and_uses_configured_model(monkeypatch):
    from app.extractors.ocr import GeminiOCRProvider
    from google import genai
    calls = []
    class Client:
        def __init__(self, **kwargs):
            self.models = self
        def __enter__(self):
            return self
        def __exit__(self, *args):
            pass
        def generate_content(self, **kwargs):
            calls.append(kwargs)
            return SimpleNamespace(text=r'$$v=\frac{s}{t}$$')
    monkeypatch.setattr(genai, 'Client', Client)
    monkeypatch.setattr(settings, 'OCR_MODEL', 'configured-vision-model')
    result = GeminiOCRProvider().recognize(b'image-bytes')
    assert result.preserves_math
    assert result.text == r'$$v=\frac{s}{t}$$'
    assert calls[0]['model'] == 'configured-vision-model'
    assert calls[0]['contents'][1].inline_data.data == b'image-bytes'


def test_force_ocr_can_recover_layout_even_when_text_is_long():
    ocr = StubOCR('Bảng và công thức được đọc lại từ ảnh trang.')
    result = DocumentProcessor(ocr=ocr).process('layout.pdf', pdf('Photosynthesis uses carbon dioxide and water and light to produce glucose in the leaves.'), ocr_mode='ocr')
    assert len(ocr.calls) == 1 and result['text'] == ocr.text


def test_native_docx_keeps_table_header_rows_and_source_section_together():
    from docx import Document
    document = Document()
    document.add_paragraph('Số liệu chuyển động')
    table = document.add_table(rows=2, cols=2)
    table.cell(0, 0).text, table.cell(0, 1).text = 'Thời gian', 'Quãng đường'
    table.cell(1, 0).text, table.cell(1, 1).text = '20 s', '100 m'
    output = BytesIO(); document.save(output)
    result = DocumentProcessor(ocr=None).process('data.docx', output.getvalue())
    tables = [b for b in result['parsed_data']['blocks'] if b['content_type'] == 'table']
    assert len(tables) == 1
    assert 'Thời gian' in tables[0]['text'] and '100 m' in tables[0]['text']
    chunks = chunk_document({'id': 12, 'filename': 'data.docx', 'file_type': 'docx', 'parsed_data': result['parsed_data'], 'extracted_text': result['text']}, chunk_size=10)
    table_chunk = next(c for c in chunks if c['content_type'] == 'table')
    assert table_chunk['metadata']['source_section'] == 'Bảng 1'
    assert table_chunk['page_number'] is None
