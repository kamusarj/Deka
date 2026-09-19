"""Phase 2 - Part 1: upload tài liệu + trích xuất (không AI).

Fixture PDF/DOCX/XLSX được sinh ngay trong test (in-memory) để không phải
commit file nhị phân.
"""

from io import BytesIO
from types import SimpleNamespace

import pytest
from fastapi import HTTPException
from app.core.config import settings
from app.extractors import (
    UnsupportedDocumentError,
    extract_docx,
    extract_pdf,
    extract_xlsx,
    resolve_file_type,
)
from app.main import app
from app.services import document_service as document_service_module
from app.services.document_service import DocumentService
from fastapi.testclient import TestClient
from sqlalchemy.exc import SQLAlchemyError


class CommitFailingSession:
    def __init__(self, document=None):
        self.document = document
        self.rollback_called = False
        self.deleted = None

    def __enter__(self):
        return self

    def __exit__(self, *_args):
        return False

    def get_bind(self):
        return SimpleNamespace(dialect=SimpleNamespace(name="postgresql"))

    def add(self, _document):
        return None

    def get(self, _model, _document_id):
        return self.document

    def query(self, _model):
        return self

    def filter(self, *_conditions):
        return self

    def first(self):
        return self.document

    def delete(self, document):
        self.deleted = document

    def commit(self):
        raise SQLAlchemyError("forced commit failure")

    def rollback(self):
        self.rollback_called = True


class CommitSucceedingSession(CommitFailingSession):
    def commit(self):
        return None


def make_pdf_bytes(lines: list[str]) -> bytes:
    from reportlab.pdfgen import canvas

    buffer = BytesIO()
    pdf = canvas.Canvas(buffer)
    y = 800
    for line in lines:
        pdf.drawString(72, y, line)
        y -= 20
    pdf.save()
    return buffer.getvalue()


def test_upload_rejects_file_larger_than_configured_limit(auth_context, monkeypatch):
    monkeypatch.setattr(settings, "MAX_UPLOAD_SIZE_BYTES", 10)

    response = TestClient(app).post(
        "/api/documents/upload",
        files={"file": ("oversized.pdf", b"x" * 11, "application/pdf")},
        headers=auth_context["headers"],
    )

    assert response.status_code == 413
    assert "Tệp vượt quá giới hạn tải lên" in response.json()["detail"]


@pytest.mark.parametrize("extension", ["pdf", "docx", "xlsx"])
@pytest.mark.parametrize("data", [b"", b"not a valid document"])
def test_invalid_document_returns_client_error_without_persistence(auth_context, monkeypatch, tmp_path, extension, data):
    from app.api.routes.documents import document_service
    monkeypatch.setattr(document_service, "upload_dir", tmp_path)
    before = document_service.list_documents(actor=auth_context["user"])
    response = TestClient(app, raise_server_exceptions=False).post(
        "/api/documents/upload",
        files={"file": (f"broken.{extension}", data, "application/octet-stream")},
        headers=auth_context["headers"],
    )
    assert response.status_code == 400, response.text
    assert "detail" in response.json() and "error_type" not in response.json()
    assert list(tmp_path.iterdir()) == []
    assert document_service.list_documents(actor=auth_context["user"]) == before


@pytest.mark.parametrize("extension", ["docx", "xlsx"])
def test_office_archive_without_document_part_is_a_client_error(auth_context, extension):
    from zipfile import ZipFile
    buffer = BytesIO()
    with ZipFile(buffer, "w") as archive:
        archive.writestr("[Content_Types].xml", '<Types xmlns="http://schemas.openxmlformats.org/package/2006/content-types"/>')
    response = TestClient(app, raise_server_exceptions=False).post(
        "/api/documents/upload", files={"file": (f"incomplete.{extension}", buffer.getvalue())},
        headers=auth_context["headers"],
    )
    assert response.status_code == 400, response.text


def test_password_protected_pdf_explains_how_to_retry(auth_context):
    from pypdf import PdfWriter
    writer = PdfWriter()
    writer.add_blank_page(width=200, height=200)
    writer.encrypt("test-document-password")
    buffer = BytesIO(); writer.write(buffer)
    response = TestClient(app, raise_server_exceptions=False).post(
        "/api/documents/upload", files={"file": ("protected.pdf", buffer.getvalue())},
        headers=auth_context["headers"],
    )
    assert response.status_code == 400
    assert "mật khẩu" in response.json()["detail"]


def test_unexpected_extractor_failure_is_not_misreported_as_bad_user_input(monkeypatch):
    from app.extractors import document_extractors
    def fail(_data):
        raise RuntimeError("Unexpected internal failure")
    monkeypatch.setitem(document_extractors._EXTRACTORS, "pdf", fail)
    with pytest.raises(RuntimeError, match="Unexpected internal failure"):
        document_extractors.extract_document("test.pdf", b"%PDF-1.7\ncontent")


def make_docx_bytes(paragraphs: list[str], table: list[list[str]] | None = None) -> bytes:
    from docx import Document

    document = Document()
    for paragraph in paragraphs:
        document.add_paragraph(paragraph)
    if table:
        docx_table = document.add_table(rows=len(table), cols=len(table[0]))
        for row_index, row in enumerate(table):
            for col_index, value in enumerate(row):
                docx_table.cell(row_index, col_index).text = value
    buffer = BytesIO()
    document.save(buffer)
    return buffer.getvalue()


def make_xlsx_bytes(rows: list[list[str]]) -> bytes:
    from openpyxl import Workbook

    workbook = Workbook()
    worksheet = workbook.active
    worksheet.title = "Ma trận"
    for row in rows:
        worksheet.append(row)
    buffer = BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


# ── Extractors thuần ────────────────────────────────────────────


def test_extract_pdf_returns_text_and_page_count():
    data = make_pdf_bytes(["Phan ung hoa hoc", "Mol va ti khoi chat khi"])
    result = extract_pdf(data)
    assert "Phan ung hoa hoc" in result["text"]
    assert result["parsed_data"]["page_count"] == 1
    assert result["parsed_data"]["page_spans"][0]["page"] == 1


def test_extract_docx_returns_paragraphs_and_table():
    data = make_docx_bytes(
        ["Phản ứng hóa học", "Dung dịch và nồng độ"],
        table=[["Chủ đề", "Số tiết"], ["Mol", "12"]],
    )
    result = extract_docx(data)
    assert "Phản ứng hóa học" in result["text"]
    assert "Chủ đề | Số tiết" in result["text"]
    assert result["parsed_data"]["paragraph_count"] == 2
    assert result["parsed_data"]["tables"][0][1] == ["Mol", "12"]
    assert result["parsed_data"]["block_spans"][0]["paragraph"] == 1


def test_extract_xlsx_returns_sheet_rows():
    data = make_xlsx_bytes([["Chủ đề", "Mức độ"], ["Phản ứng hóa học", "Nhận biết"]])
    result = extract_xlsx(data)
    assert "Phản ứng hóa học" in result["text"]
    sheet = result["parsed_data"]["sheets"][0]
    assert sheet["name"] == "Ma trận"
    assert sheet["rows"][1] == ["Phản ứng hóa học", "Nhận biết"]
    assert result["parsed_data"]["sheet_spans"][0]["sheet"] == "Ma trận"


def test_resolve_file_type_rejects_unsupported():
    assert resolve_file_type("a.pdf") == "pdf"
    assert resolve_file_type("b.DOCX") == "docx"
    try:
        resolve_file_type("c.txt")
    except UnsupportedDocumentError:
        pass
    else:  # pragma: no cover
        raise AssertionError("Phải raise UnsupportedDocumentError cho .txt")


# ── Service roundtrip ───────────────────────────────────────────


def test_document_service_roundtrip_upload_list_get_delete():
    service = DocumentService()
    saved = service.save_document(
        filename="dethi.docx",
        data=make_docx_bytes(["Nội dung tài liệu giáo viên"]),
        grade=8,
    )
    assert saved.id is not None
    assert saved.file_type == "docx"
    assert saved.grade == 8
    assert "Nội dung tài liệu giáo viên" in saved.extracted_text

    listed = service.list_documents(grade=8)
    assert any(doc.id == saved.id for doc in listed)

    fetched = service.get_document(saved.id)
    assert fetched.id == saved.id
    assert fetched.extracted_text == saved.extracted_text

    deleted = service.delete_document(saved.id)
    assert deleted.deleted is True

    # Sau khi xóa thì không còn truy vấn được.
    remaining = service.list_documents()
    assert all(doc.id != saved.id for doc in remaining)


def test_save_document_removes_new_file_when_database_commit_fails(
    tmp_path,
    monkeypatch,
):
    service = DocumentService()
    session = CommitFailingSession()
    service.upload_dir = tmp_path
    service.SessionLocal = lambda: session
    monkeypatch.setattr(
        document_service_module.DocumentProcessor,
        "process",
        lambda _self, _filename, _data, **_kwargs: {
            "file_type": "pdf",
            "text": "Nội dung",
            "parsed_data": {},
        },
    )

    with pytest.raises(SQLAlchemyError, match="forced commit failure"):
        service.save_document(filename="failure.pdf", data=b"document")

    assert session.rollback_called is True
    assert list(tmp_path.iterdir()) == []


def test_delete_document_preserves_file_when_database_commit_fails(tmp_path):
    stored_path = tmp_path / "keep-on-failure.pdf"
    stored_path.write_bytes(b"document")
    document = SimpleNamespace(id=42, stored_filename=stored_path.name, school_id=None, sharing_scope="private")
    session = CommitFailingSession(document)
    service = DocumentService()
    service.upload_dir = tmp_path
    service.SessionLocal = lambda: session

    with pytest.raises(SQLAlchemyError, match="forced commit failure"):
        service.delete_document(document.id)

    assert session.rollback_called is True
    assert session.deleted is document
    assert stored_path.read_bytes() == b"document"


@pytest.mark.parametrize("stored_filename", ["../../outside.pdf", "..\\outside.pdf"])
def test_delete_document_rejects_stored_path_traversal(tmp_path, stored_filename):
    outside_path = tmp_path.parent / "outside.pdf"
    outside_path.write_bytes(b"outside")
    document = SimpleNamespace(id=43, stored_filename=stored_filename)
    session = CommitSucceedingSession(document)
    service = DocumentService()
    service.upload_dir = tmp_path
    service.SessionLocal = lambda: session

    with pytest.raises(HTTPException, match="Tên file lưu trữ không hợp lệ"):
        service.delete_document(document.id)

    assert session.deleted is None
    assert outside_path.read_bytes() == b"outside"


# ── API layer ───────────────────────────────────────────────────


def test_upload_endpoint_roundtrip_and_rejects_bad_type(auth_context):
    client = TestClient(app)

    upload = client.post(
        "/api/documents/upload",
        files={"file": ("matran.xlsx", make_xlsx_bytes([["Câu", "Điểm"], ["1", "0.25"]]),
                        "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")},
        data={"grade": "9"},
        headers=auth_context["headers"],
    )
    assert upload.status_code == 200, upload.text
    body = upload.json()
    assert body["file_type"] == "xlsx"
    assert body["grade"] == 9
    document_id = body["id"]

    listed = client.get("/api/documents", headers=auth_context["headers"])
    assert listed.status_code == 200
    assert any(doc["id"] == document_id for doc in listed.json())

    bad = client.post(
        "/api/documents/upload",
        files={"file": ("note.txt", b"khong ho tro", "text/plain")},
        headers=auth_context["headers"],
    )
    assert bad.status_code == 400

    deleted = client.delete(f"/api/documents/{document_id}", headers=auth_context["headers"])
    assert deleted.status_code == 200

    missing = client.get(f"/api/documents/{document_id}", headers=auth_context["headers"])
    assert missing.status_code == 404
